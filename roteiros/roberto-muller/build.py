import sys, html
from docx import Document
from docx.shared import Pt, RGBColor
RED=(0xC0,0x00,0x00); BLUE=(0x1F,0x4E,0xB4); BLACK=(0,0,0)
def parse(files):
    out=[]
    for f in files:
        for p in open(f).read().split('\n\n'):
            p=p.strip()
            if not p: continue
            c=BLACK
            if p.startswith('R: '): c,p=RED,p[3:]
            elif p.startswith('A: '): c,p=BLUE,p[3:]
            out.append((c,p))
    return out
def build(title, files, base):
    paras=parse(files)
    d=Document(); st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(12)
    d.add_heading(title,1)
    for c,p in paras:
        r=d.add_paragraph().add_run(p); r.font.color.rgb=RGBColor(*c)
    d.save(base+'.docx')
    hx=lambda c:'#%02x%02x%02x'%c
    body='\n'.join(f'<p style="color:{hx(c)}">{html.escape(p)}</p>' for c,p in paras)
    open(base+'.html','w').write(f'''<!DOCTYPE html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(title)}</title>
<style>body{{background:#fff;color:#000;font-family:Georgia,serif;font-size:18px;line-height:1.6;max-width:760px;margin:0 auto;padding:24px 16px}}h1{{font-size:22px}}.leg{{font-size:14px;color:#555;border-bottom:1px solid #ddd;padding-bottom:12px}}</style></head>
<body><h1>{html.escape(title)}</h1><p class="leg"><span style="color:{hx(RED)}">■ Vermelho</span>: hook e microganchos · <span style="color:{hx(BLUE)}">■ Azul</span>: revelação do 1.º gancho e 2.º grande gancho · ■ Preto: narração</p>
{body}</body></html>''')
    print(base, sum(len(p.split()) for _,p in paras),'palavras')
if __name__=='__main__':
    build(sys.argv[1], sys.argv[3:], sys.argv[2])
